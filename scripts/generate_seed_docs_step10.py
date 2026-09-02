"""
Gera os documentos seed PDF/DOCX para o Passo 10 (parsers reais).

Conteúdo novo (não duplica os tópicos 01-08 do corpus TXT/HTML), cobrindo
temas de passos futuros do runbook — serve tanto para testar os parsers
reais quanto para popular perguntas novas no golden dataset.

Rodar com: python scripts/generate_seed_docs_step10.py
"""

import sys
from pathlib import Path

sys.path.insert(0, ".")

CORPUS_DIR = Path("data/seed_corpus")

# ---------------------------------------------------------------------------
# Conteúdo
# ---------------------------------------------------------------------------

DOC_PDF_1 = {
    "filename": "09_fast_full_path_routing.pdf",
    "title": "Fast Path e Full Path: Routing Suave",
    "sections": [
        (None, [
            "Nem toda query precisa do pipeline completo de retrieval. Rodar reranker "
            "cross-encoder e múltiplos canais de busca para uma pergunta trivial "
            "desperdiça latência e custo sem ganho de qualidade perceptível.",
        ]),
        ("O problema do custo uniforme", [
            "Sem routing, toda query paga o mesmo preço: dense search, BM25, fusão RRF "
            "e reranker, mesmo quando a resposta está obviamente em um único chunk "
            "muito específico. Isso é especialmente caro quando o reranker está ativo, "
            "já que um cross-encoder pode ser dezenas de vezes mais lento que a busca inicial.",
        ]),
        ("Dois caminhos, não dois sistemas", [
            "O fast path usa dense mais BM25 com fusao RRF uniforme, sem reranker. "
            "O full path adiciona o reranker cross-encoder sobre os candidatos do "
            "fast path. Nenhum canal de busca é eliminado em nenhum dos dois "
            "caminhos — apenas o reranker é opcional.",
        ]),
        ("Estimando complexidade sem LLM", [
            "A classificação de complexidade usa heurística leve, sem chamada de LLM: "
            "numero de entidades na query, presença de termos comparativos como "
            "'diferença entre' ou 'comparado a', e comprimento total da pergunta. "
            "Chamar um LLM só para decidir qual caminho seguir anularia o ganho de "
            "latência que o fast path deveria trazer.",
        ]),
        ("Threshold de decisão", [
            "O valor de referência adotado nesta plataforma é threshold = 0.6 na "
            "escala de complexidade estimada: abaixo disso, fast path; a partir "
            "disso, full path. Esse valor é ajustado empiricamente comparando "
            "latência média e Recall@5 do golden dataset nos dois caminhos.",
        ]),
    ],
}

DOC_PDF_2 = {
    "filename": "10_reindexacao_incremental.pdf",
    "title": "Reindexação Incremental: Grafo de Dependências",
    "sections": [
        (None, [
            "Reprocessar o corpus inteiro toda vez que um documento muda é inviável "
            "em escala. Um documento editado deve disparar apenas o recálculo dos "
            "chunks, embeddings e índices derivados dele — não do corpus inteiro.",
        ]),
        ("Detectando mudança com hash", [
            "Cada Document ganha um content_hash calculado com SHA-256 sobre o "
            "conteúdo bruto do arquivo no momento do parsing. Antes de reprocessar, "
            "o sistema compara o hash novo com o hash salvo da última indexação: "
            "se forem iguais, nada precisa ser refeito.",
        ]),
        ("O grafo de dependências", [
            "A relação document_id para chunk_ids para embedding_ids é modelada "
            "como um grafo de dependências simples. Quando um documento muda, "
            "o sistema navega o grafo a partir do document_id afetado e invalida "
            "somente os nós downstream — os chunks daquele documento e os "
            "embeddings correspondentes no Qdrant e no Postgres.",
        ]),
        ("Fluxo de invalidação", [
            "Ao detectar hash diferente: 1) remove os chunks antigos daquele "
            "document_id nos índices dense e BM25; 2) roda o parser e o chunker "
            "novamente sobre o conteúdo atualizado; 3) reindexa apenas os chunks "
            "novos. Documentos não tocados nunca entram nesse fluxo.",
        ]),
    ],
}

DOC_DOCX_1 = {
    "filename": "11_query_decomposition_hyde.docx",
    "title": "Query Decomposition e HyDE",
    "sections": [
        (None, [
            "Perguntas compostas e perguntas curtas e ambíguas são dois problemas "
            "diferentes que o retrieval híbrido básico não resolve bem sozinho. "
            "Query Decomposition ataca o primeiro; HyDE ataca o segundo.",
        ]),
        ("O problema das perguntas compostas", [
            "Uma pergunta como 'compare X e Y e explique a diferença' embute duas "
            "ou mais sub-perguntas. Buscar a pergunta composta inteira de uma vez "
            "tende a recuperar chunks genéricos que mencionam X ou Y isoladamente, "
            "raramente os dois no mesmo chunk.",
        ]),
        ("Como funciona a decomposição", [
            "A função decompose_query usa um LLM com prompt restrito para retornar "
            "apenas a lista de sub-perguntas, sem explicação adicional. Cada "
            "sub-pergunta roda o retrieval híbrido separadamente, e os resultados "
            "de todas elas são fundidos antes do reranking.",
        ]),
        ("HyDE: Hypothetical Document Embeddings", [
            "Quando a pergunta do usuário é curta e usa vocabulário muito diferente "
            "do texto-fonte, buscar pelo embedding da própria pergunta funciona mal. "
            "HyDE gera uma resposta hipotética à pergunta, mesmo que não seja "
            "precisa, e busca pelo embedding dessa resposta hipotética — que tende "
            "a se parecer mais com o texto-fonte do que a pergunta original.",
        ]),
        ("Quando usar cada técnica", [
            "Query Decomposition entra no full path para perguntas classificadas "
            "como compostas pelo router. HyDE é reservado para queries curtas e "
            "ambíguas identificadas pelo mesmo router — as duas técnicas podem "
            "coexistir, mas atacam sintomas diferentes e raramente são acionadas "
            "na mesma query.",
        ]),
    ],
}

DOC_DOCX_2 = {
    "filename": "12_migracao_embedding.docx",
    "title": "Migração de Modelo de Embedding",
    "sections": [
        (None, [
            "Trocar o modelo de embedding em produção é arriscado se feito sem um "
            "processo repetível: reindexar tudo do zero é caro, e uma migração mal "
            "feita pode derrubar o Recall@5 silenciosamente.",
        ]),
        ("Dual-write em vez de substituição", [
            "A migração cria uma segunda collection no Qdrant, em um namespace "
            "novo, para o modelo de embedding candidato. O corpus é reprocessado "
            "com o modelo novo e gravado em paralelo, sem apagar a collection "
            "antiga — as duas coexistem durante o período de avaliação.",
        ]),
        ("Avaliação lado a lado", [
            "O golden dataset roda contra as duas collections, e o Recall@5 e a "
            "Precision@5 de cada uma são comparados diretamente. Só depois dessa "
            "comparação é que a decisão de trocar o modelo padrão é tomada.",
        ]),
        ("Cutover seguro", [
            "Se o modelo novo for igual ou melhor, o EMBEDDING_MODEL padrão no "
            "config é trocado, e a collection antiga é mantida por um período "
            "antes de ser apagada — garantindo um caminho de rollback caso surjam "
            "problemas não capturados pelo golden dataset em produção.",
        ]),
    ],
}


# ---------------------------------------------------------------------------
# Geração de PDF (pymupdf)
# ---------------------------------------------------------------------------

def _wrap_text(text: str, max_chars: int = 95) -> list:
    words = text.split()
    lines, current = [], ""
    for w in words:
        candidate = f"{current} {w}".strip()
        if len(candidate) > max_chars:
            lines.append(current)
            current = w
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def write_pdf(doc_spec: dict) -> Path:
    import fitz

    out_path = CORPUS_DIR / doc_spec["filename"]
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)  # A4
    pdf.set_metadata({"title": doc_spec["title"]})

    margin_x, y = 56, 60
    page_h = 842
    bottom_margin = 60

    def new_page_if_needed(extra_h: float):
        nonlocal page, y
        if y + extra_h > page_h - bottom_margin:
            page = pdf.new_page(width=595, height=842)
            y = 60

    # Título do documento (fonte bem maior, só na primeira página)
    title_lines = _wrap_text(doc_spec["title"], max_chars=45)
    for line in title_lines:
        new_page_if_needed(26)
        page.insert_text((margin_x, y), line, fontsize=20, fontname="helv")
        y += 26
    y += 14

    for heading, paragraphs in doc_spec["sections"]:
        if heading:
            new_page_if_needed(22)
            page.insert_text((margin_x, y), heading, fontsize=14, fontname="helv")
            y += 26
        for para in paragraphs:
            for line in _wrap_text(para, max_chars=95):
                new_page_if_needed(16)
                page.insert_text((margin_x, y), line, fontsize=11, fontname="helv")
                y += 16
            y += 10
        y += 6

    pdf.save(str(out_path))
    pdf.close()
    return out_path


# ---------------------------------------------------------------------------
# Geração de DOCX (python-docx)
# ---------------------------------------------------------------------------

def write_docx(doc_spec: dict) -> Path:
    import docx

    out_path = CORPUS_DIR / doc_spec["filename"]
    d = docx.Document()
    d.core_properties.title = doc_spec["title"]
    d.add_heading(doc_spec["title"], level=0)

    for heading, paragraphs in doc_spec["sections"]:
        if heading:
            d.add_heading(heading, level=1)
        for para in paragraphs:
            d.add_paragraph(para)

    d.save(str(out_path))
    return out_path


def main() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)

    p1 = write_pdf(DOC_PDF_1)
    print(f"[OK] {p1}")
    p2 = write_pdf(DOC_PDF_2)
    print(f"[OK] {p2}")
    p3 = write_docx(DOC_DOCX_1)
    print(f"[OK] {p3}")
    p4 = write_docx(DOC_DOCX_2)
    print(f"[OK] {p4}")


if __name__ == "__main__":
    main()
