"""
Parser de DOCX — Passo 10 + Tabelas Estruturadas (Passo 15).

Usa python-docx. Diferente do PDF (que precisa inferir heading por tamanho
de fonte), o DOCX já expõe estilos de parágrafo ("Heading 1", "Heading 2",
…) diretamente — parsing mais direto e confiável.

Decisões de design:
- Parágrafos com estilo "Heading N" (ou "Título N", nome em pt-BR do Word)
  viram Section com level=N.
- Parágrafos de lista (estilo contendo "List") viram Paragraph prefixado
  com "• ", mesmo padrão do parser HTML (Passo 2).
- Tabelas (doc.tables) viram Table, com a primeira linha como headers —
  mesmo padrão do parser HTML.
- DOCX não tem conceito nativo de "página" (páginas dependem de renderização
  do Word) — não geramos elementos Page aqui, diferente do PDF.
- content_hash: SHA-256 dos bytes brutos do .docx.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import List, Optional

from src.parsing.canonical_model import Document, Image, Paragraph, Section, Table
from src.parsing.image_parser import extract_images_from_docx
from src.parsing.table_extractor import build_table_content


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


_HEADING_STYLE_RE = re.compile(r"^(heading|t[ií]tulo)\s*(\d+)$", re.IGNORECASE)


def _heading_level(style_name: str) -> Optional[int]:
    """Retorna o nível (1-9) se o estilo for um heading, senão None."""
    m = _HEADING_STYLE_RE.match(style_name.strip())
    if not m:
        return None
    try:
        return int(m.group(2))
    except ValueError:
        return None


def parse_docx(
    path: str | Path,
    tenant_id: Optional[str] = None,
    source_uri: Optional[str] = None,
    extract_images: bool = True,
    upload_storage: bool = True,
) -> Document:
    """
    Lê um arquivo .docx e produz um Document no Canonical Document Model.

    Fluxo:
    1. Percorre `document.paragraphs` na ordem do documento.
    2. Estilo "Heading N" / "Título N" -> Section(level=N).
    3. Estilo contendo "List" -> Paragraph com prefixo "• ".
    4. Demais parágrafos não-vazios -> Paragraph.
    5. Tabelas (`document.tables`) -> Table, anexadas ao final.
    6. Imagens embutidas -> Image (Passo 19), com OCR e descrição visual.
    """
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx não instalado. Execute: pip install python-docx")

    path = Path(path)
    raw_bytes = path.read_bytes()
    source = source_uri or str(path)

    word_doc = docx.Document(str(path))

    core_title = (word_doc.core_properties.title or "").strip()

    doc = Document(
        source_uri=source,
        title=core_title or None,
        tenant_id=tenant_id,
        content_hash=_sha256(raw_bytes),
        parser_version="docx-python-docx-0.1",
        language="pt",
        metadata={"original_filename": path.name, "parser": "parse_docx"},
    )

    title_set = bool(doc.title)

    for para in word_doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ""
        level = _heading_level(style_name)

        if level is not None:
            if not title_set:
                doc.title = text
                title_set = True
                continue

            sec = Section(
                title=text,
                level=level,
                content="",
                metadata={"docx_style": style_name},
            )
            doc.elements.append(sec)

        elif "list" in style_name.lower():
            p = Paragraph(content=f"• {text}", metadata={"docx_style": style_name})
            doc.elements.append(p)

        else:
            p = Paragraph(content=text, metadata={"docx_style": style_name})
            doc.elements.append(p)

    # --- Tabelas ---
    for t_idx, table in enumerate(word_doc.tables):
        rows_text: List[List[str]] = [
            [cell.text.strip() for cell in row.cells] for row in table.rows
        ]
        if not rows_text:
            continue
        headers, rows = rows_text[0], rows_text[1:]

        tbl = Table(
            headers=headers,
            rows=rows,
            metadata={"table_index": t_idx, "section_title": ""},
        )
        # Passo 15: gera content com markdown + sumário em linguagem natural
        tbl.content = build_table_content(tbl, section_title=None)
        doc.elements.append(tbl)

    # --- Imagens (Passo 19) ---
    if extract_images:
        images = extract_images_from_docx(
            word_doc,
            upload_storage=upload_storage,
        )
        for img in images:
            doc.elements.append(img)

    return doc
